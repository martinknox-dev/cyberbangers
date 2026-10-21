"""Phase 1 acceptance tests: contracts, fuzzy ingestion, deterministic simulator."""
import io
import json
import re
from datetime import date, datetime, timedelta

import openpyxl
import polars as pl
import pytest
from pydantic import ValidationError

from core.contracts import (
    ACCOUNT_COLUMNS,
    ACCOUNT_METADATA_SCHEMA,
    COMPLAINT_COLUMNS,
    COMPLAINT_SCHEMA,
    IFSC_PATTERN,
    MANDATORY_BY_TABLE,
    MANDATORY_COLUMNS,
    TRANSACTION_COLUMNS,
    TRANSACTION_SCHEMA,
    AccountMetadata,
    ColumnMapping,
    Complaint,
    IngestionReport,
    Transaction,
)

VALID_TXN_ROW: dict[str, object] = {
    "txn_id": "UTR2024060100001",
    "sender_account": "2900000001",
    "sender_ifsc": "SBIN0001234",
    "receiver_account": "2900000002",
    "receiver_ifsc": "SBIN0009876",
    "amount": 15000.0,
    "timestamp": datetime(2024, 6, 1, 10, 30, 0),
    "payment_mode": "UPI",
    "device_id": "D000123",
    "ip_address": "10.20.30.40",
}


class TestContracts:
    def test_valid_transaction_row_validates(self) -> None:
        txn = Transaction.model_validate(VALID_TXN_ROW)
        assert txn.amount == 15000.0
        assert txn.payment_mode == "UPI"

    def test_transaction_rejects_negative_amount(self) -> None:
        with pytest.raises(ValidationError):
            Transaction.model_validate({**VALID_TXN_ROW, "amount": -1.0})

    def test_transaction_rejects_bad_payment_mode(self) -> None:
        with pytest.raises(ValidationError):
            Transaction.model_validate({**VALID_TXN_ROW, "payment_mode": "VISA-KEEP"})

    def test_transaction_rejects_empty_txn_id(self) -> None:
        with pytest.raises(ValidationError):
            Transaction.model_validate({**VALID_TXN_ROW, "txn_id": ""})

    def test_account_metadata_constraints(self) -> None:
        ok = AccountMetadata(
            account_no="2900000001",
            ifsc="SBIN0001234",
            account_opening_date=date(2023, 1, 15),
            branch_code="BR001",
            mobile_no="9178451234",
            imei="351104510123456",
            kyc_pincode="800013",
            account_type="PMJDY",
        )
        assert ok.account_type == "PMJDY"
        with pytest.raises(ValidationError):
            AccountMetadata(
                account_no="1", ifsc="BAD", branch_code="B", account_type="JOINT"  # type: ignore[arg-type]
            )

    def test_complaint_roundtrip_json(self) -> None:
        c = Complaint(
            ack_no="ACK20240001",
            victim_account="2900000007",
            suspect_account="5500000001",
            fraud_category="AEPS_CASHOUT",
            reported_timestamp=datetime(2024, 6, 2, 9, 0, 0),
            disputed_amount=32000.0,
            disputed_utr="UTR2024060100001",
            state="Bihar",
        )
        dumped = json.loads(c.model_dump_json())
        assert dumped["disputed_amount"] == 32000.0
        assert dumped["disputed_utr"] == "UTR2024060100001"

    def test_complaint_disputed_utr_optional(self) -> None:
        with_utr = Complaint(
            ack_no="ACK20240002",
            victim_account="2900000007",
            fraud_category="UPI_FRAUD",
            reported_timestamp=datetime(2024, 6, 2, 9, 0, 0),
            disputed_amount=1500.0,
            state="Bihar",
        )
        assert with_utr.disputed_utr is None   # nullable, links only when known

    def test_report_defaults_hold_dataframe(self) -> None:
        report = IngestionReport()
        assert report.rows_total == 0 and report.rows_accepted == 0
        assert report.rejected_rows.is_empty()

    def test_ifsc_pattern_is_strict(self) -> None:
        assert IFSC_PATTERN.fullmatch("SBIN0001234") is not None
        assert IFSC_PATTERN.fullmatch("sbin0001234") is None
        assert IFSC_PATTERN.fullmatch("SBIN00123") is None

    def test_transaction_schema_is_string_first(self) -> None:
        assert list(TRANSACTION_COLUMNS) == [
            "txn_id", "sender_account", "sender_ifsc", "receiver_account",
            "receiver_ifsc", "amount", "timestamp", "payment_mode",
            "device_id", "ip_address",
        ]
        assert TRANSACTION_SCHEMA["amount"] == pl.Float64
        assert TRANSACTION_SCHEMA["timestamp"] == pl.Datetime("us")
        for c in TRANSACTION_SCHEMA:
            if c not in ("amount", "timestamp"):
                assert TRANSACTION_SCHEMA[c] == pl.String, c

    def test_account_and_complaint_schemas(self) -> None:
        assert [ACCOUNT_METADATA_SCHEMA[c] for c in ("account_opening_date", "account_type")] == [
            pl.Date, pl.String,
        ]
        assert list(COMPLAINT_SCHEMA) == [
            "ack_no", "victim_account", "suspect_account", "fraud_category",
            "reported_timestamp", "disputed_amount", "disputed_utr", "state",
        ]
        assert list(COMPLAINT_COLUMNS) == list(COMPLAINT_SCHEMA)
        assert list(ACCOUNT_COLUMNS) == list(ACCOUNT_METADATA_SCHEMA)
        assert [COMPLAINT_SCHEMA[c] for c in ("reported_timestamp", "disputed_amount")] == [
            pl.Datetime("us"), pl.Float64,
        ]

    def test_mandatory_columns_subset_of_schema(self) -> None:
        assert MANDATORY_COLUMNS == (
            "txn_id", "sender_account", "receiver_account",
            "amount", "timestamp", "payment_mode",
        )
        assert set(MANDATORY_COLUMNS) <= set(TRANSACTION_SCHEMA)
        assert MANDATORY_BY_TABLE == {
            "transactions": MANDATORY_COLUMNS,
            "complaints": (
                "ack_no", "victim_account", "fraud_category",
                "reported_timestamp", "disputed_amount", "state",
            ),
            "accounts": ("account_no", "ifsc", "branch_code", "account_type"),
        }
        assert set(MANDATORY_BY_TABLE["complaints"]) <= set(COMPLAINT_SCHEMA)
        assert set(MANDATORY_BY_TABLE["accounts"]) <= set(ACCOUNT_METADATA_SCHEMA)

    def test_column_mapping_confidence_bounds(self) -> None:
        m = ColumnMapping(raw_name="Payee", canonical_name="receiver_account",
                          matched_synonym="payee", confidence=1.0)
        assert m.confidence == 1.0
        with pytest.raises(ValidationError):
            ColumnMapping(raw_name="x", canonical_name=None, confidence=1.5)


from ingestion.mapper import DatasetIngester, map_headers, normalize_header

STYLE_A_HEADERS: list[str] = [   # core-banking CSV export (bank style, underscored codes)
    "UTR", "Debit_Acc", "Fr_IFSC", "PAYEE", "To_IFSC", "Amount_INR",
    "Txn_Time", "Mode", "Device", "IP",
]
STYLE_B_HEADERS: list[str] = [   # branch XLSX sheet (spaced title-case)
    "Transaction Ref", "Originator Acct", "Remitter IFSC", "To Account",
    "Beneficiary Branch IFSC", "Txn Amount", "Txn DateTime", "Pay Method",
    "Device No", "Customer IP",
]
STYLE_C_HEADERS: list[str] = [   # complaint-portal dump (mixed snake + CamelCase)
    "txn ref id", "sender account", "Sender IFSC Code", "beneficiary_account",
    "Receiver IFSC Code", "amount", "datetime_of_txn", "channel",
    "imei_device", "IP",
]

STYLE_A_EXPECTED: dict[str, str] = {
    "UTR": "txn_id", "Debit_Acc": "sender_account", "Fr_IFSC": "sender_ifsc",
    "PAYEE": "receiver_account", "To_IFSC": "receiver_ifsc",
    "Amount_INR": "amount", "Txn_Time": "timestamp", "Mode": "payment_mode",
    "Device": "device_id", "IP": "ip_address",
}
STYLE_B_EXPECTED: dict[str, str] = {
    "Transaction Ref": "txn_id", "Originator Acct": "sender_account",
    "Remitter IFSC": "sender_ifsc", "To Account": "receiver_account",
    "Beneficiary Branch IFSC": "receiver_ifsc", "Txn Amount": "amount",
    "Txn DateTime": "timestamp", "Pay Method": "payment_mode",
    "Device No": "device_id", "Customer IP": "ip_address",
}
STYLE_C_EXPECTED: dict[str, str] = {
    "txn ref id": "txn_id", "sender account": "sender_account",
    "Sender IFSC Code": "sender_ifsc", "beneficiary_account": "receiver_account",
    "Receiver IFSC Code": "receiver_ifsc", "amount": "amount",
    "datetime_of_txn": "timestamp", "channel": "payment_mode",
    "imei_device": "device_id", "IP": "ip_address",
}

COMPLAINT_HEADERS: list[str] = [   # NCRP-portal style dump (report-officer wording)
    "Complaint Ack No", "Victim Acc No", "Suspect Acc No", "Fraud Category",
    "Reported On", "Cheated Amount", "Disputed UTR Ref", "State",
]
ACCOUNT_HEADERS: list[str] = [    # core-banking customer master sheet
    "Account Number", "IFSC Code", "Date of Opening", "Home Branch Code",
    "Mobile Number", "Device IMEI", "KYC Pincode", "Acct Type",
]
COMPLAINT_EXPECTED: dict[str, str] = {
    "Complaint Ack No": "ack_no", "Victim Acc No": "victim_account",
    "Suspect Acc No": "suspect_account", "Fraud Category": "fraud_category",
    "Reported On": "reported_timestamp", "Cheated Amount": "disputed_amount",
    "Disputed UTR Ref": "disputed_utr", "State": "state",
}
ACCOUNT_EXPECTED: dict[str, str] = {
    "Account Number": "account_no", "IFSC Code": "ifsc",
    "Date of Opening": "account_opening_date", "Home Branch Code": "branch_code",
    "Mobile Number": "mobile_no", "Device IMEI": "imei",
    "KYC Pincode": "kyc_pincode", "Acct Type": "account_type",
}
MULTI_TABLE_EXACT_CONFS: dict[tuple[str, str], float] = {
    # (table, raw header): probe-computed confidence
    ("complaints", "Complaint Ack No"): 0.8966,
    ("complaints", "Disputed UTR Ref"): 0.8571,
    ("accounts", "Date of Opening"): 0.8889,
    ("accounts", "Home Branch Code"): 0.8148,
}


class TestHeaderMapping:
    def test_normalize_header(self) -> None:
        assert normalize_header("Beneficiary-Acc  No#2") == "beneficiary acc no 2"
        assert normalize_header("  UTR ") == "utr"

    def test_style_a_full_coverage(self) -> None:
        mapping = map_headers(STYLE_A_HEADERS)
        for raw, canonical in STYLE_A_EXPECTED.items():
            assert mapping[raw].canonical_name == canonical, raw
            assert mapping[raw].confidence >= 0.80, (raw, mapping[raw].confidence)

    def test_style_b_full_coverage(self) -> None:
        mapping = map_headers(STYLE_B_HEADERS)
        for raw, canonical in STYLE_B_EXPECTED.items():
            assert mapping[raw].canonical_name == canonical, raw
            assert mapping[raw].confidence >= 0.80, raw

    def test_style_c_full_coverage_with_exact_confidences(self) -> None:
        mapping = map_headers(STYLE_C_HEADERS)
        exact_confs: dict[str, float] = {
            "txn ref id": 0.824, "sender account": 1.0, "Sender IFSC Code": 0.815,
            "beneficiary_account": 1.0, "Receiver IFSC Code": 0.839,
            "amount": 1.0, "datetime_of_txn": 0.889, "channel": 1.0,
            "imei_device": 0.80, "IP": 1.0,
        }
        for raw, canonical in STYLE_C_EXPECTED.items():
            assert mapping[raw].canonical_name == canonical, raw
            assert abs(mapping[raw].confidence - exact_confs[raw]) <= 1e-3, raw

    def test_complaint_and_account_style_coverage_with_exact_confs(self) -> None:
        m_c = map_headers(COMPLAINT_HEADERS, table="complaints")
        for raw, canonical in COMPLAINT_EXPECTED.items():
            assert m_c[raw].canonical_name == canonical, raw
            assert m_c[raw].confidence >= 0.80, raw
        m_a = map_headers(ACCOUNT_HEADERS, table="accounts")
        for raw, canonical in ACCOUNT_EXPECTED.items():
            assert m_a[raw].canonical_name == canonical, raw
            assert m_a[raw].confidence >= 0.80, raw
        # exact-confidence checks (the marginal headers per table):
        for raw in ("Complaint Ack No", "Disputed UTR Ref"):
            assert abs(m_c[raw].confidence - MULTI_TABLE_EXACT_CONFS[("complaints", raw)]) <= 1e-3, raw
        for raw in ("Date of Opening", "Home Branch Code"):
            assert abs(m_a[raw].confidence - MULTI_TABLE_EXACT_CONFS[("accounts", raw)]) <= 1e-3, raw

    def test_default_call_maps_transactions_only(self) -> None:
        # Without `table=`, complaint/account-specific headers must NOT claim
        # transaction canonicals, and transaction headers must not claim
        # complaint canonicals (tables are siloed by the synonym registry).
        tx_style = map_headers(["Cheated Amount", "Victim Acc No"])
        assert tx_style["Cheated Amount"].canonical_name is None
        assert tx_style["Victim Acc No"].canonical_name is None
        cx_style = map_headers(["Amount_INR", "Txn_Time"], table="complaints")
        assert cx_style["Amount_INR"].canonical_name is None
        assert cx_style["Txn_Time"].canonical_name is None

    def test_noise_headers_unmapped(self) -> None:
        mapping = map_headers(
            ["Narration", "Remarks", "Sl No", "Balance", "SenderIFSC", "ReceiverBranchIFSC"]
        )
        for raw, entry in mapping.items():
            assert entry.canonical_name is None, raw
            assert entry.confidence < 0.80, raw

    def test_below_threshold_is_dropped(self) -> None:
        mapping = map_headers(["Narration"], confidence_threshold=0.55)
        assert mapping["Narration"].canonical_name is None  # best score 52.2 < 55

    def test_canonical_collision_exact_tie_keeps_first(self) -> None:
        mapping = map_headers(["To Account", "PAYEE"])  # both hit receiver_account @ 100
        assert mapping["To Account"].canonical_name == "receiver_account"
        assert mapping["PAYEE"].canonical_name is None     # tie -> first-seen wins
        assert mapping["PAYEE"].confidence == 1.0          # loser keeps its own score

    def test_canonical_collision_higher_confidence_wins(self) -> None:
        # "Credit_Acc_No" scores ~0.741 against credit_account; "To Account" 1.0.
        mapping = map_headers(["Credit_Acc_No", "To Account"])
        assert mapping["To Account"].canonical_name == "receiver_account"
        assert mapping["Credit_Acc_No"].canonical_name is None
        assert mapping["Credit_Acc_No"].confidence < 0.80


def csv_bytes(rows: list[str]) -> bytes:
    header = ",".join(STYLE_A_HEADERS) + "\n"
    body = "\n".join(rows)
    return (header + body + ("\n" if body else "")).encode("utf-8")


def style_a_xlsx(header: list[str], rows: list[list[object]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for c, name in enumerate(header, start=1):
        ws.cell(column=c, row=1, value=name)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(column=c, row=r, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestIngestHappyPath:
    def test_csv_path_all_three_header_styles_accepted(self, tmp_path) -> None:
        style_rows: dict[str, list[str]] = {
            "a": [
                "UTR_A1,2900000001,SBIN0001234,2900000002,SBIN0009876,15000.50,2024-06-01 10:30:00,UPI,D000123,10.20.30.40",
            ],
            "b": [
                "UTR_B1,2900000001,SBIN0001234,2900000002,SBIN0009876,2499.00,01/06/2024 11:45:30,IMPS,D000124,10.20.30.41",
            ],
            "c": [
                "UTR_C1,2900000001,SBIN0001234,2900000002,SBIN0009876,777.75,2024-06-01T12:00:00,NEFT,D000125,10.20.30.42",
            ],
        }
        headers_by: dict[str, list[str]] = {
            "a": STYLE_A_HEADERS, "b": STYLE_B_HEADERS, "c": STYLE_C_HEADERS,
        }
        for key, rows in style_rows.items():
            path = tmp_path / f"style_{key}.csv"
            path.write_bytes(
                (",".join(headers_by[key]) + "\n" + rows[0] + "\n").encode("utf-8")
            )
            df, report = DatasetIngester().ingest(str(path))
            assert df.height == 1
            assert report.rows_total == 1
            assert report.rows_accepted == 1 and report.rows_rejected == 0
            assert list(df.columns) == list(TRANSACTION_COLUMNS)
            assert df.dtypes == [TRANSACTION_SCHEMA[c] for c in TRANSACTION_COLUMNS]
            assert df["txn_id"].item() == f"UTR_{key.upper()}1"

    def test_dd_mm_precedes_mm_dd(self, tmp_path) -> None:
        rows = [
            "UTR_D1,2900000001,SBIN0001234,2900000002,SBIN0009876,500.00,05/06/2024 09:15:00,UPI,,",
            "UTR_D2,2900000001,SBIN0001234,2900000002,SBIN0009876,500.00,12/31/2024 09:15:00,UPI,,",
            "UTR_D3,2900000001,SBIN0001234,2900000002,SBIN0009876,500.00,05/06/2024 09:15:52 PM,UPI,,",
        ]
        path = tmp_path / "dd_mm.csv"
        path.write_bytes((",".join(STYLE_A_HEADERS) + "\n" + "\n".join(rows) + "\n").encode())
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 0
        stamps = {r["txn_id"]: r["timestamp"] for r in df.to_dicts()}
        assert stamps["UTR_D1"] == datetime(2024, 6, 5, 9, 15, 0)
        assert stamps["UTR_D2"] == datetime(2024, 12, 31, 9, 15, 0)
        assert stamps["UTR_D3"] == datetime(2024, 6, 5, 21, 15, 52)

    def test_unix_timestamp_seconds_and_ms(self, tmp_path) -> None:
        rows = [
            "UTR_E1,1,SBIN0000001,2,SBIN0000002,10,1717200000,UPI,,",
            "UTR_E2,1,SBIN0000001,2,SBIN0000002,10,1717200002000,IMPS,,",
        ]
        path = tmp_path / "epoch.csv"
        path.write_bytes((",".join(STYLE_A_HEADERS) + "\n" + "\n".join(rows) + "\n").encode())
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 0
        assert sorted(df["timestamp"].to_list()) == [
            datetime(2024, 6, 1, 0, 0, 0), datetime(2024, 6, 1, 0, 0, 2),
        ]

    def test_xlsx_bytes_with_native_cells(self) -> None:
        payload = style_a_xlsx(
            STYLE_A_HEADERS,
            [[
                "UTR_X1", "2900000001", "sbin0001234", "2900000002",
                "SBIN0009876", 100.5, datetime(2024, 6, 1, 3, 4, 5), "BHIM UPI",
                "D000001", "10.0.0.1",
            ]],
        )
        df, report = DatasetIngester().ingest(payload)   # bytes -> xlsx via PK magic sniff
        assert report.rows_accepted == 1
        assert df["amount"].item() == pytest.approx(100.5)
        assert df["timestamp"].item() == datetime(2024, 6, 1, 3, 4, 5)
        assert df["payment_mode"].item() == "UPI"
        assert df["sender_ifsc"].item() == "SBIN0001234"   # lowercased input upcased + validated

    def test_bytes_csv_via_sniff(self) -> None:
        payload = csv_bytes([
            "UTR_F1,1,SBIN0000001,2,SBIN0000002,25.00,2024-06-01 10:00:00,CASH,dev1,1.2.3.4",
        ])
        df, report = DatasetIngester().ingest(payload)
        assert report.rows_accepted == 1
        assert df["payment_mode"].item() == "CASH"

    def test_explicit_reader_override_on_suffixless_path(self, tmp_path) -> None:
        path = tmp_path / "only"
        path.write_bytes(csv_bytes([
            "UTR_G1,1,SBIN0000001,2,SBIN0000002,25.00,2024-06-01 10:00:00,CASH,dev2,1.2.3.5",
        ]))
        as_csv, report = DatasetIngester().ingest_csv(str(path))
        assert as_csv.height == 1 and report.rows_accepted == 1
        with pytest.raises(openpyxl.utils.exceptions.InvalidFileException):
            DatasetIngester().ingest_xlsx(str(path))

    def test_mapping_report_records_confidence(self, tmp_path) -> None:
        path = tmp_path / "m.csv"
        path.write_bytes(csv_bytes([
            "UTR_H1,1,SBIN0000001,2,SBIN0000002,25.00,2024-06-01 10:00:00,CASH,dev3,1.2.3.6",
        ]))
        _df, report = DatasetIngester().ingest(str(path))
        m = report.column_mapping["Debit_Acc"]
        assert m.canonical_name == "sender_account"
        assert m.matched_synonym == "debit_account"
        assert m.confidence >= 0.80
        assert report.column_mapping["UTR"].canonical_name == "txn_id"


def complaint_csv_bytes(rows: list[str]) -> bytes:
    header = ",".join(COMPLAINT_HEADERS) + "\n"
    return (header + "\n".join(rows) + "\n").encode("utf-8")


def account_xlsx_bytes(rows: list[list[object]]) -> bytes:
    return style_a_xlsx(ACCOUNT_HEADERS, rows)


class TestIngestMultiTableHappy:
    def test_complaints_csv_happy_path(self, tmp_path) -> None:
        path = tmp_path / "ncrp.csv"
        path.write_bytes(complaint_csv_bytes([
            'ACK20240001,2900000007,2990100000,AEPS_CASHOUT,02/06/2024 09:15:00,"32,000.50",UTR2024060100001,Bihar',
        ]))
        df, report = DatasetIngester().ingest_complaints(str(path))
        assert df.height == 1
        assert list(df.columns) == list(COMPLAINT_COLUMNS)
        assert df["ack_no"].item() == "ACK20240001"
        assert df["disputed_utr"].item() == "UTR2024060100001"
        # DD/MM precedence + comma-junk amount survive the shared pipeline:
        assert df["reported_timestamp"].item() == datetime(2024, 6, 2, 9, 15, 0)
        assert df["disputed_amount"].item() == pytest.approx(32000.5)
        assert df["suspect_account"].item() == "2990100000"
        assert report.rows_accepted == 1 and report.rows_rejected == 0

    def test_complaints_headers_do_not_claim_transaction_columns(self, tmp_path) -> None:
        path = tmp_path / "ncrp.csv"
        path.write_bytes(complaint_csv_bytes([
            "ACK20240002,2900000007,2990100001,UPI_FRAUD,02/06/2024 09:15:00,1500,UTR2024060100002,Delhi",
        ]))
        df, report = DatasetIngester().ingest_complaints(str(path))
        assert df.height == 1 and report.rows_accepted == 1

    def test_accounts_xlsx_native_date_cell(self) -> None:
        payload = account_xlsx_bytes([[
            "2900000007", "sbin0001234", date(2023, 1, 15), "BR001",
            9178451234, "351104510123456", "800013", "savings",
        ]])
        df, report = DatasetIngester().ingest_accounts(payload)   # PK magic sniff
        assert df.height == 1
        assert list(df.columns) == list(ACCOUNT_COLUMNS)
        # openpyxl-native datetime cell -> String "2023-01-15 00:00:00.000000" -> Date:
        assert df["account_opening_date"].item() == date(2023, 1, 15)
        assert df["account_type"].item() == "SAVINGS"     # lowercase alias upcased
        assert df["ifsc"].item() == "SBIN0001234"         # lowercase input upcased
        assert df["mobile_no"].item() == "9178451234"     # numeric cell cast to String
        assert df["imei"].item() == "351104510123456"
        assert report.rows_accepted == 1

    def test_payment_mode_alias_expansion(self, tmp_path) -> None:
        rows = [
            "UTR_N1,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:00:00,RTGS,,",
            "UTR_N2,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:05:00,ATM WITHDRAWAL,,",
            "UTR_N3,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:10:00,MICRO ATM,,",
            "UTR_N4,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:15:00,UPI/P2M,,",
            "UTR_N5,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:20:00,UPI/P2P,,",
        ]
        path = tmp_path / "modes.csv"
        path.write_bytes((",".join(STYLE_A_HEADERS) + "\n" + "\n".join(rows) + "\n").encode())
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 0
        assert df["payment_mode"].to_list() == ["RTGS", "ATM", "AEPS", "UPI", "UPI"]


class TestIngestRejections:
    ROW = (
        "UTR_R1,2900000001,SBIN0001234,2900000002,SBIN0009876,100.00,"
        "2024-06-01 10:30:00,UPI,D000001,10.0.0.9"
    )

    def _ingest_rows(self, tmp_path, rows: list[str]):
        path = tmp_path / "reject.csv"
        path.write_bytes(csv_bytes(rows))
        return DatasetIngester().ingest(str(path))

    def test_negative_amount_rejected(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            self.ROW,
            "UTR_R2,2900000001,SBIN0001234,2900000002,SBIN0009876,-5.00,2024-06-01 10:31:00,UPI,,",
        ])
        assert report.rows_total == 2
        assert report.rows_rejected == 1 and df.height == 1
        assert report.failure_counts == {"negative_amount": 1}
        assert report.rejected_rows["UTR"].item() == "UTR_R2"
        assert report.rejected_rows["rejection_reason"].item() == "negative_amount"

    def test_unparseable_timestamp_rejected(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            "UTR_R3,1,SBIN0000001,2,SBIN0000002,10.00,NOT-A-DATE,UPI,,",
        ])
        assert report.failure_counts == {"unparseable_timestamp": 1}
        assert df.height == 0

    def test_unparseable_amount_rejected(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            "UTR_R4,1,SBIN0000001,2,SBIN0000002,ten,2024-06-01 10:30:00,UPI,,",
        ])
        assert report.failure_counts == {"unparseable_amount": 1}
        assert df.height == 0

    def test_missing_mandatory_fields(self, tmp_path) -> None:
        rows = [
            ",1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:30:00,UPI,,",              # missing:txn_id
            "UTR_M2,,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:30:00,UPI,,",         # missing:sender_account
            "UTR_M3,1,SBIN0000001,,SBIN0000002,10,2024-06-01 10:30:00,UPI,,",         # missing:receiver_account
            "UTR_M4,1,SBIN0000001,2,SBIN0000002,,2024-06-01 10:30:00,UPI,,",          # missing:amount
            "UTR_M5,1,SBIN0000001,2,SBIN0000002,10,,UPI,,",                            # missing:timestamp
            "UTR_M6,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:30:00,,,1.1.1.1",     # missing:payment_mode
        ]
        df, report = self._ingest_rows(tmp_path, rows)
        assert df.height == 0
        assert report.rows_rejected == 6
        assert report.failure_counts == {
            "missing:txn_id": 1, "missing:sender_account": 1,
            "missing:receiver_account": 1, "missing:amount": 1,
            "missing:timestamp": 1, "missing:payment_mode": 1,
        }

    def test_duplicate_utr_keep_first(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            self.ROW,
            self.ROW.replace("100.00", "200.00"),
        ])
        assert report.failure_counts == {"duplicate_utr": 1}
        assert df.height == 1 and df["amount"].item() == 100.0
        assert report.rejected_rows["UTR"].item() == "UTR_R1"

    def test_invalid_ifsc_rejected(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            "UTR_R5,1,SBIN1234,2,SBIN0000002,10,2024-06-01 10:30:00,UPI,,",
        ])
        assert report.failure_counts == {"invalid_ifsc": 1}
        assert df.height == 0

    def test_unsupported_payment_mode_rejected(self, tmp_path) -> None:
        _df, report = self._ingest_rows(tmp_path, [
            "UTR_R6,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:30:00,CHEQUE,,",
        ])
        assert report.failure_counts == {"unsupported_payment_mode": 1}

    def test_priority_missing_beats_all(self, tmp_path) -> None:
        _df, report = self._ingest_rows(tmp_path, [
            "UTR_R7,,SBIN0000001,2,SBIN0000002,-10,NOT-A-DATE,ZZZ,,",
        ])
        assert report.failure_counts == {"missing:sender_account": 1}

    def test_mixed_reasons_aggregate(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            "UTR_R8,1,SBIN0000001,2,SBIN0000002,-1,2024-06-01 10:30:00,UPI,,",          # negative_amount
            "UTR_R9,1,SBIN0000001,2,SBIN0000002,10,NOPE,UPI,,",                          # unparseable_timestamp
        ])
        assert report.failure_counts == {"negative_amount": 1, "unparseable_timestamp": 1}
        assert report.rows_accepted == 0 and df.height == 0
        assert set(report.rejected_rows.columns) == set(STYLE_A_HEADERS) | {"rejection_reason"}

    def test_zero_amount_accepted(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            "UTR_Z1,1,SBIN0000001,2,SBIN0000002,0,2024-06-01 10:30:00,UPI,,",
        ])
        assert report.rows_rejected == 0 and df.height == 1
        assert df["amount"].item() == 0.0

    def test_header_only_csv_returns_empty_report(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [])
        assert df.height == 0 and df.width == len(TRANSACTION_COLUMNS)
        assert report.rows_total == 0 and report.rows_accepted == 0
        assert report.rows_rejected == 0 and report.failure_counts == {}
        assert report.column_mapping["UTR"].canonical_name == "txn_id"
        assert report.rejected_rows.is_empty()

    def test_ragged_csv_rows_do_not_crash(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            self.ROW + ",EXTRA,CELLS",          # reader truncates extras; row parses
            self.ROW.replace("UTR_R1", "UTR_R10"),
        ])
        assert report.rows_total == 2
        assert report.rows_rejected == 0
        assert df.height == 2

    def test_quoted_timestamp_cell_still_parses(self, tmp_path) -> None:
        df, report = self._ingest_rows(tmp_path, [
            'UTR_CX1,1,SBIN0000001,2,SBIN0000002,10,"2024-06-01 10:30:00",UPI,,',
        ])
        assert report.rows_rejected == 0
        assert df["timestamp"].item() == datetime(2024, 6, 1, 10, 30, 0)


class TestIngestMultiTableRejections:
    def test_complaint_missing_ack_duplicate_ack(self, tmp_path) -> None:
        rows = [
            # empty ack -> missing:ack_no (also: suspect + disputed UTR optional, fine)
            ",2900000007,2990100000,AEPS_CASHOUT,02/06/2024 09:15:00,32000,UTR2024060100001,Bihar",
            "ACK20240005,2900000008,2990100001,UPI_FRAUD,02/06/2024 10:15:00,1500,UTR2024060100002,Delhi",
            "ACK20240005,2900000008,2990100001,UPI_FRAUD,02/06/2024 10:20:00,1500,UTR2024060100002,Delhi",  # dup ack
        ]
        path = tmp_path / "ncrp_reject.csv"
        path.write_bytes(complaint_csv_bytes(rows))
        df, report = DatasetIngester().ingest_complaints(str(path))
        assert report.rows_total == 3
        assert df.height == 1
        assert df["ack_no"].item() == "ACK20240005"
        assert report.failure_counts == {"duplicate_ack_no": 1, "missing:ack_no": 1}
        assert report.rejected_rows["rejection_reason"].to_list()[0] == "missing:ack_no"

    def test_complaint_missing_mandatory_field(self, tmp_path) -> None:
        rows = [
            # empty state -> missing:state; amount/timestamp fine
            "ACK20240006,2900000007,2990100002,AEPS_CASHOUT,02/06/2024 09:15:00,9000,UTR2024060100003,",
        ]
        path = tmp_path / "ncrp_state.csv"
        path.write_bytes(complaint_csv_bytes(rows))
        df, report = DatasetIngester().ingest_complaints(str(path))
        assert report.failure_counts == {"missing:state": 1}
        assert df.height == 0

    def test_accounts_rejections(self, tmp_path) -> None:
        rows = [
            # ifsc malformed -> invalid_ifsc
            "2994000000,SBIN00,15/03/2024,BR042,9812000000,351104510900001,800013,CURRENT",
            # unsupported account type -> unsupported_account_type
            "2994000001,SBIN0001234,15/03/2024,BR042,9812000001,351104510900002,800013,JOINT",
            # pincode not digits -> invalid_kyc_pincode
            "2994000002,SBIN0001234,15/03/2024,BR042,9812000002,351104510900003,12AB56,CURRENT",
            # mobile too short -> invalid_mobile_no (5 digits)
            "2994000003,SBIN0001234,15/03/2024,BR042,98123,351104510900004,800013,CURRENT",
            # clean row -> accepted, plus duplicate account_no below
            "2994000004,SBIN0001234,2024-03-15,BR042,9812000004,351104510900005,800013,SAVINGS",
            "2994000004,SBIN0001234,2024-03-20,BR042,9812000005,351104510900006,800013,SAVINGS",  # dup account_no
        ]
        path = tmp_path / "accounts_reject.csv"
        path.write_bytes((",".join(ACCOUNT_HEADERS) + "\n" + "\n".join(rows) + "\n").encode())
        df, report = DatasetIngester().ingest_accounts(str(path))
        assert report.rows_total == 6
        assert report.rows_accepted == 1 and df.height == 1
        assert df["account_no"].item() == "2994000004"
        assert df["account_opening_date"].item() == date(2024, 3, 15)   # keep-first
        assert report.failure_counts == {
            "duplicate_account_no": 1, "invalid_ifsc": 1, "invalid_kyc_pincode": 1,
            "invalid_mobile_no": 1, "unsupported_account_type": 1,
        }

    def test_accounts_missing_mandatory(self, tmp_path) -> None:
        rows = [
            # empty branch code -> missing:branch_code
            "2994000005,SBIN0001234,16/03/2024,,9812000006,351104510900007,800013,CURRENT",
        ]
        path = tmp_path / "accounts_missing.csv"
        path.write_bytes((",".join(ACCOUNT_HEADERS) + "\n" + "\n".join(rows) + "\n").encode())
        df, report = DatasetIngester().ingest_accounts(str(path))
        assert report.failure_counts == {"missing:branch_code": 1}
        assert df.height == 0


from simulator.generator import (
    LABELED_TRANSACTION_SCHEMA,
    SCHEME_AEPS_TERMINAL_CASH_OUT,
    SCHEME_BRANCH_COLLUSION,
    SCHEME_CIRCULAR_WASH,
    SCHEME_FAN_OUT_SMURFING,
    SCHEME_RAPID_PASS_THROUGH,
    SCHEME_TXN_COUNTS,
    SHARED_BRANCH_IMEI,
    TERMINAL_MERCHANT_ACCOUNT,
    SimulationResult,
    SyndicateDataSimulator,
)


class TestSimulatorBackground:
    def test_default_shape_and_validation(self) -> None:
        result = SyndicateDataSimulator(seed=42).generate(background_transactions=5_000)
        t = result.transactions
        # height = background count + emitted labeled rows (0 in Task 6; 249 after Task 7)
        assert t.height == 5_000 + int(t["is_fraud"].sum())
        with pytest.raises(ValueError):
            SyndicateDataSimulator().generate(background_transactions=4_999)
        with pytest.raises(ValueError):
            SyndicateDataSimulator().generate(background_transactions=20_001)

    def test_schema_conformance_and_defaults(self) -> None:
        result = SyndicateDataSimulator(seed=42).generate()
        assert result.transactions.columns == list(LABELED_TRANSACTION_SCHEMA)
        assert result.accounts.columns == list(ACCOUNT_METADATA_SCHEMA)
        assert result.accounts.height == 1274
        assert result.transactions.dtypes == [
            LABELED_TRANSACTION_SCHEMA[c] for c in result.transactions.columns
        ]

    def test_background_all_unlabeled(self) -> None:
        result = SyndicateDataSimulator().generate(background_transactions=6_000)
        bg = result.transactions.filter(~pl.col("is_fraud"))
        assert bg.height == 6_000
        assert bg["fraud_scheme"].null_count() == bg.height

    def test_txn_ids_unique_and_background_prefixed(self) -> None:
        result = SyndicateDataSimulator().generate()
        ids = result.transactions["txn_id"]
        assert ids.n_unique() == ids.len()
        bg_ids = result.transactions.filter(~pl.col("is_fraud"))["txn_id"]
        assert bg_ids.str.starts_with("BG").all()

    def test_all_accounts_referenceable(self) -> None:
        result = SyndicateDataSimulator().generate()
        accounts = set(result.accounts["account_no"].to_list())
        used = (
            set(result.transactions["sender_account"].to_list())
            | set(result.transactions["receiver_account"].to_list())
        )
        assert used <= accounts

    def test_transactions_sorted(self) -> None:
        result = SyndicateDataSimulator().generate()
        expected = result.transactions.sort(["timestamp", "txn_id"])
        assert result.transactions.equals(expected)

    def test_deterministic_same_seed(self) -> None:
        a = SyndicateDataSimulator(seed=42).generate()
        b = SyndicateDataSimulator(seed=42).generate()
        assert a.transactions.equals(b.transactions)
        assert a.accounts.equals(b.accounts)

    def test_different_seed_changes_stream(self) -> None:
        a = SyndicateDataSimulator(seed=42).generate(background_transactions=5_000)
        b = SyndicateDataSimulator(seed=7).generate(background_transactions=5_000)
        assert not a.transactions.equals(b.transactions)


_TERMINAL_ACCOUNT = TERMINAL_MERCHANT_ACCOUNT   # single source of truth lives in the generator
_SHARED_IMEI = SHARED_BRANCH_IMEI


def accounts_set(sim: SimulationResult) -> set[str]:
    return set(sim.accounts["account_no"].to_list())


class TestSchemes:
    @pytest.fixture()
    def sim(self) -> SimulationResult:
        return SyndicateDataSimulator(seed=42).generate(background_transactions=5_000)

    def test_labeled_counts_exact(self, sim) -> None:
        got = {
            row["fraud_scheme"]: row["len"]
            for row in sim.transactions.group_by("fraud_scheme").len().to_dicts()
        }
        labeled = {k: v for k, v in got.items() if k is not None}
        assert labeled == SCHEME_TXN_COUNTS

    def test_total_rows_match_background_plus_schemes(self, sim) -> None:
        assert sim.transactions.height == 5_000 + sum(SCHEME_TXN_COUNTS.values())

    def test_labels_integrity(self, sim) -> None:
        t = sim.transactions
        matches = t.select(
            (pl.col("is_fraud") == pl.col("fraud_scheme").is_not_null()).all()
        ).item()
        assert bool(matches)

    def test_no_orphaned_or_letter_coded_accounts(self, sim) -> None:
        # Decision 13: every emitted sender/receiver — background and all five
        # schemes — resolves to a master row; ids are numeric only.
        t = sim.transactions
        used = set(t["sender_account"].to_list()) | set(t["receiver_account"].to_list())
        assert used <= accounts_set(sim)
        assert not any(re.search(r"[A-Za-z]", account) for account in used)
        assert sim.accounts.height == 1_274          # 1000 base + 273 scheme + 1 merchant

    def test_rapid_pass_through_drain_window_hub_age(self, sim) -> None:
        t = sim.transactions.filter(pl.col("fraud_scheme") == SCHEME_RAPID_PASS_THROUGH)
        assert t.height == 80
        for e in range(20):
            rows = t.filter(
                pl.col("txn_id").str.starts_with(f"S1-{e:03d}-")
            ).sort("txn_id")
            inflow = rows.filter(pl.col("txn_id").str.ends_with("-0"))["amount"].item()
            hops = rows.filter(~pl.col("txn_id").str.ends_with("-0"))
            drain = hops["amount"].sum() / inflow
            assert drain >= 0.90, (e, drain)
            window = hops["timestamp"].max() - rows["timestamp"].min()
            assert window <= timedelta(minutes=15)
            assert hops.height == 3
            # roles derived structurally: the inflow leg's receiver IS the hub
            hub = t.filter(pl.col("txn_id") == f"S1-{e:03d}-0")["receiver_account"].item()
            hub_sent = rows.filter(pl.col("sender_account") == hub)
            assert hub_sent["amount"].sum() <= inflow  # hub can't spend its inflow twice
            # attribute truth for Phase 2 (Decision 13): hub account is aged, low-tier
            hub_row = sim.accounts.filter(pl.col("account_no") == hub)
            assert hub_row.height == 1
            opened = hub_row["account_opening_date"].item()
            assert opened <= date(2024, 6, 1) - timedelta(days=180)
            assert hub_row["account_type"].item() in ("PMJDY", "SAVINGS")

    def test_fan_out_smurfing_micro_amounts(self, sim) -> None:
        t = sim.transactions.filter(pl.col("fraud_scheme") == SCHEME_FAN_OUT_SMURFING)
        for e in range(10):
            episode = t.filter(pl.col("txn_id").str.starts_with(f"S2-{e:03d}-"))
            inflow = episode.filter(pl.col("txn_id").str.ends_with("-0"))
            assert inflow.height == 1
            coll = inflow["receiver_account"].item()          # structural collector id
            micro = episode.filter(~pl.col("txn_id").str.ends_with("-0"))
            assert micro.height == 12
            assert (micro["sender_account"] == coll).all()    # collector emits all micro-sends
            assert micro["receiver_account"].n_unique() == 12
            assert (micro["amount"] < 10_000).all()

    def test_aeps_terminal_cash_out_common_point(self, sim) -> None:
        t = sim.transactions.filter(pl.col("fraud_scheme") == SCHEME_AEPS_TERMINAL_CASH_OUT)
        inbound = t.filter(pl.col("receiver_account") == _TERMINAL_ACCOUNT)
        assert inbound.height == 3
        assert (inbound["payment_mode"] == "AEPS").all()
        assert (inbound["device_id"] == "AEPSDEV000001").all()
        for trail in range(3):
            # role derivation: leg 0 is victim->tier1, leg 1 tier1->tier2, leg 2 tier2->terminal
            episode = t.filter(
                pl.col("txn_id").str.starts_with(f"S3-{trail:02d}-")
            ).sort("txn_id")
            senders = episode["sender_account"].to_list()
            receivers = episode["receiver_account"].to_list()
            victim, tier1, tier2 = senders[0], receivers[0], senders[2]
            assert receivers[1] == tier2
            assert tier1 not in {victim, tier2}
            actors = {victim, tier1, tier2, _TERMINAL_ACCOUNT}
            assert actors <= accounts_set(sim)

    def test_branch_collusion_accounts_and_device(self, sim) -> None:
        colluders = sim.accounts.filter(pl.col("imei") == _SHARED_IMEI)
        assert colluders.height == 6
        assert (colluders["ifsc"] == "SBIN0000421").all()
        assert (colluders["branch_code"] == "BR042").all()
        opens = colluders["account_opening_date"].to_list()
        assert max(opens) - min(opens) <= timedelta(days=2)
        for mob in colluders["mobile_no"].to_list():
            assert mob.startswith("9162")
        t = sim.transactions.filter(pl.col("fraud_scheme") == SCHEME_BRANCH_COLLUSION)
        assert t.height == 12
        assert (t["device_id"] == _SHARED_IMEI).all()
        colluder_set = set(colluders["account_no"].to_list())
        for row in t.to_dicts():
            assert row["sender_account"] in colluder_set
            assert row["receiver_account"] in colluder_set

    def test_circular_wash_closed_loop(self, sim) -> None:
        t = sim.transactions.filter(pl.col("fraud_scheme") == SCHEME_CIRCULAR_WASH)
        assert t.height == 18
        for loop in range(6):
            episode = t.filter(
                pl.col("txn_id").str.starts_with(f"S5-{loop:02d}-")   # builder pads 2 digits
            ).sort("txn_id")
            edges = episode.to_dicts()      # legs commute in txn_id order: 0,1,2
            assert edges[1]["sender_account"] == edges[0]["receiver_account"]   # A->B, B->C
            assert edges[2]["sender_account"] == edges[1]["receiver_account"]   # then C->...
            assert edges[0]["sender_account"] == edges[2]["receiver_account"]   # ...->A (closed)
            amounts = [edge["amount"] for edge in edges]
            assert amounts[0] > amounts[1] > amounts[2]      # 3% commission decay
            actors = {
                edge["sender_account"] for edge in edges
            } | {
                edge["receiver_account"] for edge in edges
            }
            assert actors <= accounts_set(sim)
            assert len(actors) == 3

    def test_invalid_background_transaction_count_rejected(self) -> None:
        with pytest.raises(ValueError):
            SyndicateDataSimulator().generate(background_transactions=1)


class TestComplaints:
    @pytest.fixture()
    def sim(self) -> SimulationResult:
        return SyndicateDataSimulator(seed=42).generate(background_transactions=5_000)

    def test_complaint_ground_truth_joins(self, sim) -> None:
        c = sim.complaints
        assert c.height == 16
        assert c.columns == list(COMPLAINT_SCHEMA)
        assert c["ack_no"].n_unique() == 16
        assert c["disputed_utr"].null_count() == 0
        named_states = set(c["state"].to_list())
        assert named_states == {"Bihar", "Delhi", "Maharashtra", "Tamil Nadu", "Jharkhand"}
        joined = c.join(
            sim.transactions.select(["txn_id", "sender_account", "receiver_account"]),
            left_on="disputed_utr", right_on="txn_id",
            how="left",
        )
        # polars drops the right join key (`txn_id`) in this join shape, so
        # matched-ness is asserted via the joined data columns:
        assert joined["sender_account"].null_count() == 0
        assert joined["receiver_account"].null_count() == 0
        # every complaint cites a real transaction, and victim/suspect are
        # exactly that disputable leg's sender/receiver (Decision 14):
        assert (joined["victim_account"] == joined["sender_account"]).all()
        assert (joined["suspect_account"] == joined["receiver_account"]).all()

    def test_complaints_per_scheme(self, sim) -> None:
        c = sim.complaints
        joined = c.join(
            sim.transactions.select(["txn_id", "fraud_scheme"]),
            left_on="disputed_utr", right_on="txn_id",
            how="left",
        )
        scheme_counts = {
            str(row["fraud_scheme"]): row["len"]
            for row in joined.group_by("fraud_scheme").len().to_dicts()
        }
        assert scheme_counts == {
            SCHEME_RAPID_PASS_THROUGH: 8,
            SCHEME_FAN_OUT_SMURFING: 5,
            SCHEME_AEPS_TERMINAL_CASH_OUT: 3,
        }


CANONICAL_HEADER = (
    "txn_id,sender_account,sender_ifsc,receiver_account,receiver_ifsc,"
    "amount,timestamp,payment_mode,device_id,ip_address\n"
)
CANONICAL_ROW = (
    "UTR_Q1,2900000001,SBIN0001234,2900000002,SBIN0009876,100.00,"
    "2024-06-01 10:30:00,UPI,D000001,10.0.0.9"
)


class TestHostileInputs:
    """No input may crash the ingester; every bad row carries one reason."""

    def _csv(self, tmp_path, header: str, rows: list[str], name: str = "h.csv"):
        path = tmp_path / name
        path.write_bytes((header + "\n".join(rows) + ("\n" if rows else "")).encode())
        return path

    def test_out_of_range_epoch_rejected_not_panicking(self, tmp_path) -> None:
        # 15-digit (microsecond-ish) and 16-digit epochs are not supported
        # units: they must be rejected, never materialised as year 7411/56385.
        for value in ("171720000000000", "1717200000000000"):
            path = self._csv(
                tmp_path,
                ",".join(STYLE_A_HEADERS) + "\n",
                [f"UTR_E,1,SBIN0000001,2,SBIN0000002,10,{value},UPI,,"],
                "epoch.csv",
            )
            df, report = DatasetIngester().ingest(str(path))
            assert report.rows_accepted == 0, value
            assert report.failure_counts == {"unparseable_timestamp": 1}, value
            assert df.height == 0

    def test_supported_epoch_units_still_accepted(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            ",".join(STYLE_A_HEADERS) + "\n",
            [
                "UTR_ES,1,SBIN0000001,2,SBIN0000002,10,1717200000,UPI,,",
                "UTR_EM,1,SBIN0000001,2,SBIN0000002,10,1717200000000,UPI,,",
            ],
            "epoch_ok.csv",
        )
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 0
        assert sorted(df["timestamp"].to_list()) == [
            datetime(2024, 6, 1, 0, 0, 0), datetime(2024, 6, 1, 0, 0, 0),
        ]

    def test_non_finite_amounts_rejected(self, tmp_path) -> None:
        for raw in ("NaN", "inf", "Infinity", "1e400"):
            path = self._csv(
                tmp_path,
                ",".join(STYLE_A_HEADERS) + "\n",
                [f"UTR_N,1,SBIN0000001,2,SBIN0000002,{raw},2024-06-01 10:30:00,UPI,,"],
                "nan.csv",
            )
            df, report = DatasetIngester().ingest(str(path))
            assert report.rows_accepted == 0, raw
            assert report.failure_counts == {"unparseable_amount": 1}, raw
            assert df.height == 0

    def test_empty_csv_never_raises(self, tmp_path) -> None:
        path = tmp_path / "empty.csv"
        path.write_bytes(b"")
        df, report = DatasetIngester().ingest(str(path))
        assert df.height == 0 and df.width == len(TRANSACTION_COLUMNS)
        assert report.rows_total == 0 and report.rows_accepted == 0
        assert report.rows_rejected == 0 and report.failure_counts == {}
        assert report.rejected_rows.is_empty()

    def test_whitespace_only_csv_never_raises(self, tmp_path) -> None:
        path = tmp_path / "blank.csv"
        path.write_bytes(b"   \n")
        df, report = DatasetIngester().ingest(str(path))
        assert df.height == 0 and df.width == len(TRANSACTION_COLUMNS)
        assert report.rows_total == 0

    def test_header_only_xlsx_never_raises(self) -> None:
        payload = style_a_xlsx(STYLE_A_HEADERS, [])
        df, report = DatasetIngester().ingest(payload)
        assert df.height == 0 and df.width == len(TRANSACTION_COLUMNS)
        assert report.rows_total == 0 and report.rows_accepted == 0
        assert report.failure_counts == {}
        assert report.column_mapping["UTR"].canonical_name == "txn_id"

    def test_completely_empty_xlsx_never_raises(self) -> None:
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        df, report = DatasetIngester().ingest(buf.getvalue())
        assert df.height == 0 and df.width == len(TRANSACTION_COLUMNS)
        assert report.rows_total == 0

    def test_audit_preserves_source_text_for_canonical_headers(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            CANONICAL_HEADER,
            [CANONICAL_ROW.replace("100.00", "-5.00")],
            "canon.csv",
        )
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 1 and df.height == 0
        rejected = report.rejected_rows
        # source text, not the normalized/typed value
        assert rejected["amount"].item() == "-5.00"
        assert rejected["timestamp"].item() == "2024-06-01 10:30:00"
        assert rejected["rejection_reason"].item() == "negative_amount"
        assert rejected.schema["amount"] == pl.String

    def test_source_column_named_rejection_reason_does_not_crash(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            ",".join([*STYLE_A_HEADERS, "rejection_reason"]) + "\n",
            [CANONICAL_ROW.replace("100.00", "-5.00") + ",SOURCE-NOTE"],
            "reason.csv",
        )
        _df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 1
        rejected = report.rejected_rows
        assert rejected["rejection_reason"].item() == "negative_amount"
        assert rejected["rejection_reason_source"].item() == "SOURCE-NOTE"

    def test_source_column_named_raw_txn_id_keeps_source_value(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            ",".join([*STYLE_A_HEADERS, "_raw_txn_id"]) + "\n",
            [CANONICAL_ROW.replace("100.00", "-5.00") + ",ZZZ-ORIGINAL"],
            "rawtxn.csv",
        )
        _df, report = DatasetIngester().ingest(str(path))
        assert report.rows_rejected == 1
        assert report.rejected_rows["_raw_txn_id"].item() == "ZZZ-ORIGINAL"

    def test_malformed_account_opening_date_rejected(self, tmp_path) -> None:
        for raw_date in ("NOT-A-DATE", "2024-02-31"):
            path = self._csv(
                tmp_path,
                ",".join(ACCOUNT_HEADERS) + "\n",
                [f"2994000010,SBIN0001234,{raw_date},BR042,9812000010,351104510900010,800013,CURRENT"],
                "acct.csv",
            )
            _df, report = DatasetIngester().ingest_accounts(str(path))
            assert report.rows_accepted == 0, raw_date
            assert report.failure_counts == {"invalid_account_opening_date": 1}, raw_date

    def test_invalid_imei_rejected(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            ",".join(ACCOUNT_HEADERS) + "\n",
            ["2994000011,SBIN0001234,15/03/2024,BR042,9812000011,12345,800013,CURRENT"],
            "imei.csv",
        )
        df, report = DatasetIngester().ingest_accounts(str(path))
        assert report.failure_counts == {"invalid_imei": 1}
        assert df.height == 0

    def test_rejection_reason_vocabulary(self) -> None:
        from ingestion.mapper import REJECTION_REASONS

        for table, reasons in REJECTION_REASONS.items():
            assert reasons, table
            assert len(set(reasons)) == len(reasons), table

    def test_rejected_row_does_not_consume_its_dedupe_key(self, tmp_path) -> None:
        # The bad row comes first and shares the UTR with the good row: the
        # good row must still be accepted (dedupe runs among survivors only).
        path = self._csv(
            tmp_path,
            ",".join(STYLE_A_HEADERS) + "\n",
            [
                "UTR_DUP,1,SBIN0000001,2,SBIN0000002,-5.00,2024-06-01 10:30:00,UPI,,",
                "UTR_DUP,1,SBIN0000001,2,SBIN0000002,10.00,2024-06-01 10:31:00,UPI,,",
            ],
            "dup.csv",
        )
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_accepted == 1 and df.height == 1
        assert df["txn_id"].item() == "UTR_DUP"
        assert report.rows_rejected == 1
        assert report.failure_counts == {"negative_amount": 1}

    def test_report_counts_partition_the_source(self, tmp_path) -> None:
        path = self._csv(
            tmp_path,
            ",".join(STYLE_A_HEADERS) + "\n",
            [
                "UTR_P1,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:30:00,UPI,,",
                "UTR_P1,1,SBIN0000001,2,SBIN0000002,10,2024-06-01 10:31:00,UPI,,",
                "UTR_P2,1,SBIN0000001,2,SBIN0000002,-1,2024-06-01 10:32:00,UPI,,",
                "UTR_P3,1,SBIN0000001,2,SBIN0000002,10,NOPE,UPI,,",
            ],
            "partition.csv",
        )
        df, report = DatasetIngester().ingest(str(path))
        assert report.rows_total == 4
        assert report.rows_total == report.rows_accepted + report.rows_rejected
        assert report.rows_accepted == df.height == 1
        assert report.failure_counts == {
            "duplicate_utr": 1, "negative_amount": 1, "unparseable_timestamp": 1,
        }
